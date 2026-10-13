"""
scripts/manual_builder/build.py
-------------------------------
Master assembly builder that imports all parts and produces:
FedSSL_TB_Engineering_Mastery_Manual.docx
"""

import sys
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure scripts package is in path
ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.manual_builder.part0_foreword import build_part0
from scripts.manual_builder.part1_architecture import build_part1
from scripts.manual_builder.part2_repo_map import build_part2
from scripts.manual_builder.part3_git import build_part3
from scripts.manual_builder.part4_trace import build_part4
from scripts.manual_builder.part5_data import build_part5
from scripts.manual_builder.part6_models import build_part6
from scripts.manual_builder.part7_federated import build_part7
from scripts.manual_builder.part8_fewshot import build_part8
from scripts.manual_builder.part9_experiments import build_part9
from scripts.manual_builder.part10_paper import build_part10
from scripts.manual_builder.part11_debugging import build_part11
from scripts.manual_builder.part12_placement import build_part12
from scripts.manual_builder.part13_study_plan import build_part13
from scripts.manual_builder.part14_viva import build_part14
from scripts.manual_builder.part15_glossary import build_part15


def setup_document():
    doc = docx.Document()
    # Configure 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header and Footer
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("FedSSL-TB Engineering Mastery Manual  |  Confidential Research Documentation")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        footer = section.footer
        p_ftr = footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Federated Self-Supervised Learning for Tuberculosis Detection  •  Page ")
        r_ftr.font.name = "Calibri"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    return doc


def main():
    print("=" * 70)
    print("Building FedSSL_TB_Engineering_Mastery_Manual.docx...")
    print("=" * 70)

    doc = setup_document()

    parts = [
        ("Part 0: Foreword & Audit Reconciliation", build_part0),
        ("Part 1: Mental Model & Architecture", build_part1),
        ("Part 2: Engineering Repository Map", build_part2),
        ("Part 3: Git & GitHub Workflow", build_part3),
        ("Part 4: Terminal-to-Result Execution Trace", build_part4),
        ("Part 5: Medical Data Pipeline", build_part5),
        ("Part 6: Model Internals & Optimization", build_part6),
        ("Part 7: Federated Optimization", build_part7),
        ("Part 8: Few-Shot Metric Learning", build_part8),
        ("Part 9: Empirical Experiments & Registry", build_part9),
        ("Part 10: IEEE Paper Preparation", build_part10),
        ("Part 11: Engineering Troubleshooting Runbook", build_part11),
        ("Part 12: Technical Placement Mastery", build_part12),
        ("Part 13: 12-Week Study Programme", build_part13),
        ("Part 14: Viva Voce 35 Q&A", build_part14),
        ("Part 15: Glossary, Source Index & References", build_part15),
    ]

    for title, builder_fn in parts:
        print(f"  -> Building {title}...")
        builder_fn(doc)

    out_root = ROOT / "FedSSL_TB_Engineering_Mastery_Manual.docx"
    doc.save(str(out_root))
    print(f"\n[SUCCESS] Saved primary manual to: {out_root}")

    # Copy to docs/ and Learning Materials/
    docs_dir = ROOT / "docs"
    docs_dir.mkdir(exist_ok=True)
    out_docs = docs_dir / "FedSSL_TB_Engineering_Mastery_Manual.docx"
    shutil.copy2(str(out_root), str(out_docs))
    print(f"[SUCCESS] Copied to: {out_docs}")

    learn_dir = ROOT.parent / "Learning Materials"
    if learn_dir.exists():
        out_learn = learn_dir / "FedSSL_TB_Engineering_Mastery_Manual.docx"
        shutil.copy2(str(out_root), str(out_learn))
        print(f"[SUCCESS] Copied to: {out_learn}")

    print("=" * 70)
    print("Generation complete! Total paragraphs:", len(doc.paragraphs), "Tables:", len(doc.tables))
    print("=" * 70)


if __name__ == "__main__":
    main()
