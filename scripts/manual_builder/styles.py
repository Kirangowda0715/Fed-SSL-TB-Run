"""
scripts/manual_builder/styles.py
--------------------------------
Professional styling engine for the FedSSL-TB Engineering Mastery Manual.
Provides typography, color palettes, callouts, tables, and code block styling.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

# Color Palette
COLOR_NAVY_DARK   = RGBColor(0x1A, 0x23, 0x4E)  # #1A234E - Title & H1
COLOR_NAVY_MID    = RGBColor(0x16, 0x50, 0x8A)  # #16508A - H2 & Subtitle
COLOR_FOREST      = RGBColor(0x2E, 0x7D, 0x32)  # #2E7D32 - H3 & Verified
COLOR_PURPLE      = RGBColor(0x5E, 0x22, 0x75)  # #5E2275 - H4
COLOR_BODY        = RGBColor(0x22, 0x22, 0x22)  # #222222 - Body text
COLOR_MUTED       = RGBColor(0x66, 0x66, 0x66)  # #666666 - Captions & notes
COLOR_ALERT_NOTE  = RGBColor(0x0D, 0x47, 0xA1)  # Blue
COLOR_ALERT_TIP   = RGBColor(0x1B, 0x5E, 0x20)  # Green
COLOR_ALERT_WARN  = RGBColor(0xB7, 0x41, 0x0E)  # Orange-rust
COLOR_ALERT_CRIT  = RGBColor(0xB7, 0x1C, 0x1C)  # Red
COLOR_EVIDENCE    = RGBColor(0x00, 0x4D, 0x40)  # Teal/Dark Cyan

FONT_BODY = "Calibri"
FONT_CODE = "Consolas"
FONT_HEADING = "Calibri"


def set_paragraph_shading(paragraph, fill_hex: str):
    """Set background color of a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def set_cell_shading(cell, fill_hex: str):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell internal padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for margin_name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{margin_name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color_hex="CCCCCC", sz="4"):
    """Apply clean subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_h1(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_HEADING
    run.font.size = Pt(20)
    run.font.color.rgb = COLOR_NAVY_DARK
    run.bold = True
    return p


def add_h2(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_HEADING
    run.font.size = Pt(15)
    run.font.color.rgb = COLOR_NAVY_MID
    run.bold = True
    return p


def add_h3(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_HEADING
    run.font.size = Pt(12.5)
    run.font.color.rgb = COLOR_FOREST
    run.bold = True
    return p


def add_h4(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = FONT_HEADING
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_PURPLE
    run.bold = True
    return p


def add_body(doc, text: str, bold_prefix: str = None, italic: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = FONT_BODY
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.color.rgb = COLOR_BODY
        r_prefix.bold = True
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_BODY
    run.italic = italic
    return p


def add_bullet(doc, text: str, bold_prefix: str = None, level: int = 0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = FONT_BODY
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.color.rgb = COLOR_BODY
        r_prefix.bold = True
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_BODY
    return p


def add_code_block(doc, code_str: str, language: str = ""):
    """Add a shaded code block with Consolas font."""
    lines = code_str.strip().split("\n")
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.right_indent = Inches(0.25)
        set_paragraph_shading(p, "F2F4F7")
        run = p.add_run(line if line else " ")
        run.font.name = FONT_CODE
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    # Spacer after code block
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)


def add_callout(doc, title: str, text: str, callout_type: str = "NOTE"):
    """
    Render a single-cell callout box with a thick left border.
    callout_type: NOTE (blue), TIP (green), WARNING (orange), EVIDENCE (teal), PROFESSOR (navy)
    """
    colors = {
        "NOTE": ("EBF3FC", "1976D2", COLOR_ALERT_NOTE),
        "TIP": ("EAF5EA", "388E3C", COLOR_ALERT_TIP),
        "WARNING": ("FDF2E9", "E65100", COLOR_ALERT_WARN),
        "CRITICAL": ("FDE8E8", "D32F2F", COLOR_ALERT_CRIT),
        "EVIDENCE": ("E0F2F1", "00796B", COLOR_EVIDENCE),
        "PROFESSOR": ("E8EAF6", "283593", COLOR_NAVY_DARK),
    }
    bg_hex, border_hex, title_rgb = colors.get(callout_type.upper(), ("F4F5F7", "757575", COLOR_BODY))

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

    # Thick left border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r_icon = p.add_run(f"[{callout_type.upper()}] {title}\n")
    r_icon.font.name = FONT_HEADING
    r_icon.font.size = Pt(10.5)
    r_icon.font.color.rgb = title_rgb
    r_icon.bold = True

    r_text = p.add_run(text)
    r_text.font.name = FONT_BODY
    r_text.font.size = Pt(10.0)
    r_text.font.color.rgb = COLOR_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, headers, rows, col_widths=None):
    """Add a professionally styled table with navy headers and alternating shading."""
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "D0D5DD")

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header_text
        set_cell_shading(cell, "1A234E")  # Dark Navy
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = FONT_HEADING
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.bold = True

    # Data Rows
    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        bg_hex = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            cell = row_cells[col_idx]
            cell.text = str(cell_value)
            set_cell_shading(cell, bg_hex)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            for run in p.runs:
                run.font.name = FONT_BODY
                run.font.size = Pt(9.5)
                run.font.color.rgb = COLOR_BODY

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)
